# -*- coding: utf-8 -*-
"""将 docs/星辰书城-数据库课程设计报告.md 转换为提交版 Word 文档。

用法：python scripts/build_report_docx.py
输出：docs/202400300104-许皓宸-数据库课程设计报告.docx
- 封面 + 可更新目录域（Word 打开后自动提示更新域）
- 章节标题走 Heading 样式（目录可识别）
- 表格、代码块、图片、粗体、行内代码均按报告排版处理
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "docs" / "星辰书城-数据库课程设计报告.md"
OUT_PATH = ROOT / "docs" / "202400300104-许皓宸-数据库课程设计报告.docx"

SONG, HEI, KAI = "宋体", "黑体", "楷体"
CODE_FONT = "Consolas"


def set_fonts(run, east=SONG, west="Times New Roman", size=12, bold=False, color=None):
    run.font.name = west
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)


def add_rich(par, text, size=12, east=SONG, west="Times New Roman", bold_all=False):
    """解析 **粗体** 与 `行内代码` 的简单行内格式。"""
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        bold = seg.startswith("**") and seg.endswith("**") and len(seg) > 4
        body = seg[2:-2] if bold else seg
        for sub in re.split(r"(`[^`]+`)", body):
            if not sub:
                continue
            if sub.startswith("`") and sub.endswith("`") and len(sub) > 2:
                r = par.add_run(sub[1:-1])
                set_fonts(r, east=east, west=CODE_FONT, size=size - 1.5, bold=bold or bold_all)
            else:
                r = par.add_run(sub)
                set_fonts(r, east=east, west=west, size=size, bold=bold or bold_all)


def body_par(doc, text, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(3)
    if indent:
        pf.first_line_indent = Pt(24)
    add_rich(p, text)
    return p


def heading(doc, text, level):
    p = doc.add_heading("", level=level)
    sizes = {1: 15, 2: 13, 3: 12}
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(8 if level == 1 else 6)
    add_rich(p, text, size=sizes[level], east=HEI, west="Arial", bold_all=True)
    for r in p.runs:  # 覆盖 Word 默认标题蓝色 → 正式报告黑色
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.left_indent = Cm(0.4)
        r = p.add_run(line if line else " ")
        set_fonts(r, east=SONG, west=CODE_FONT, size=8.5, color=(60, 60, 60))
        # 代码底色
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F5F5F5")
        p._p.get_or_add_pPr().append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j in range(ncol):
            text = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_rich(p, text, size=9, east=HEI if i == 0 else SONG, bold_all=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, alt, path):
    img = ROOT / path
    if not img.exists():
        body_par(doc, f"【缺图：{path}】", indent=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img), width=Cm(15))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(alt)
    set_fonts(r, east=SONG, size=9, color=(100, 100, 100))


def add_page_number(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = OxmlElement("w:r")
    fld_b = OxmlElement("w:fldChar"); fld_b.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_e = OxmlElement("w:fldChar"); fld_e.set(qn("w:fldCharType"), "end")
    for node in (fld_b, instr, fld_e):
        run.append(node)
    p._p.append(run)


def enable_update_fields(doc):
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


def add_toc(doc):
    p = doc.add_paragraph()  # 目录标题用普通段落，避免出现在目录自身
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("目　录")
    set_fonts(r, east=HEI, west="Arial", size=16, bold=True)
    p = doc.add_paragraph()
    run = OxmlElement("w:r")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    tip = OxmlElement("w:t")
    tip.text = "（在 Word 中右键此处选择“更新域”即可生成目录）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld, instr, sep, tip, end):
        run.append(node)
    p._p.append(run)


def build_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, "数据库课程设计报告", size=26, east=HEI, west="Arial", bold_all=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, "——网文书城系统（星辰书城）", size=16, east=KAI, bold_all=True)
    for _ in range(5):
        doc.add_paragraph()
    for line in ("软件学院", "软件工程专业　2024 级＿＿班",
                 "姓名：许皓宸", "学号：202400300104",
                 "任课教师：魏淑越", "实验教师：＿＿＿"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        add_rich(p, line, size=14, east=SONG, bold_all=True)
    doc.add_page_break()


def main():
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.8)
    add_page_number(doc)
    enable_update_fields(doc)
    build_cover(doc)
    add_toc(doc)
    doc.add_page_break()

    # 跳过 MD 中的封面与手写目录，从正文第一章开始
    start = next(i for i, l in enumerate(lines) if l.startswith("# 一、"))
    i, in_code, code_buf, table_buf = start, False, [], []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            flush_table()
            if in_code:
                add_code_block(doc, code_buf)
                code_buf = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        flush_table()
        s = line.strip()
        if not s or s == "---":
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            heading(doc, m.group(2), len(m.group(1)))
        elif re.match(r"^!\[.*\]\(.*\)$", s):
            alt, path = re.match(r"^!\[(.*)\]\((.*)\)$", s).groups()
            add_image(doc, alt, path)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_rich(p, s[2:])
        elif re.match(r"^\d+[\.、]\s*", s):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            pf.left_indent = Cm(0.75)
            add_rich(p, s)
        else:
            body_par(doc, s)
        i += 1
    flush_table()
    doc.save(OUT_PATH)
    print(f"saved: {OUT_PATH}")


if __name__ == "__main__":
    sys.exit(main())
